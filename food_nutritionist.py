from PIL import Image
from datetime import datetime
from docx import Document
from docx.shared import Inches
import os
from MainJarvis_TaskExecution import speak
import google.generativeai as genai

# Configure Gemini API


food_prompt = """
You are JARVIS, an expert AI nutritionist. Analyze the food in the provided image.

Your tasks:
- Identify each food item present.
- For each food item, provide:
  - Name
  - Estimated serving size
  - Calories (kcal)
  - Protein (g)
  - Fat (g)
  - Carbohydrates (g)

- After listing all food items, provide a summary:
  - Total Calories
  - Total Protein, Fat, and Carbohydrates
  - Macronutrient breakdown in percentages (Carbs, Fats, Proteins, Fiber, Sugar)
  - Is the meal healthy? (Yes/No)
  - Any health concerns or notes

Guidelines:
- Present food item data in bullet-point format (no tables).
- All numbers must use standard nutritional units.
- If any value is assumed, clearly mention it.
- Keep the analysis short, clear, and structured.
- Avoid markdown symbols like *, #, or tables.
"""

def analyze_food_nutrition(image_path):
    GEMINI_API_KEY = os.getenv("GEMINI_API_3")  # Replace with your actual key
    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel("gemini-1.5-pro")
    """Analyzes food items and generates a structured nutrition report."""
    if not os.path.exists(image_path):
        speak("❌ Sorry Sir, the specified image file does not exist.")
        return

    try:
        speak("Analyzing the food image sent by you.")
        img = Image.open(image_path).convert('RGB')

        response = model.generate_content([food_prompt, img])

        if response and response.text:
            report_text = response.text.strip()

            # Create Word document
            doc = Document()
            doc.add_heading('🍽️ Food Nutrition Analysis Report', 0)

            # Add image (optional)
            try:
                doc.add_picture(image_path, width=Inches(4.5))
            except:
                doc.add_paragraph("(Image could not be embedded)")

            # Add timestamp
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            doc.add_paragraph(f"📅 Date of Analysis: {timestamp}")

            # Add analysis
            doc.add_paragraph(report_text)

            # Save document
            safe_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"Food_Report_{safe_timestamp}.docx"
            report_path = os.path.join(os.getcwd(), filename)
            doc.save(report_path)

            speak("Sir, the food nutrition report has been successfully generated.")
            print(f"Report saved as {filename}")
            speak("Opening the file for you")
            os.startfile(report_path)
            speak("File opened successfully")
            return report_path
        else:
            speak("❌ Sorry Sir, I was unable to analyze the food image.")
            return

    except Exception as e:
        print(e)
        speak("❌ An error occurred while analyzing the food image.")
        return

if __name__ == "__main__":
    path = input("Enter image path: ")
    analyze_food_nutrition(path)
