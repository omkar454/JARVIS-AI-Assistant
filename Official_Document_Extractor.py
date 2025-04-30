import os
from datetime import datetime
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
import google.generativeai as genai
from MainJarvis_TaskExecution import speak

# Set your Gemini API key


# Initialize Gemini model (no chat session needed)
model = genai.GenerativeModel("gemini-1.5-pro")

# 📄 Prompt for general official documents
official_doc_prompt = '''
You are an expert at interpreting official documents. Analyze the image of the document provided to you and extract all the key details.

Your analysis should include:
- Document type (e.g., invoice, certificate, contract, memo, letter, etc.)
- Issued date or any visible date
- Name of parties involved (issuer, receiver, signatories)
- Reference numbers or document IDs
- Purpose or subject of the document
- Any key figures or monetary amounts
- Stamps, seals, or official marks
- Any additional legal or important notes

Format the output in detail with clear headings. Avoid markdown. Keep the structure clean and organized.
'''

def analyze_official_document(image_path):
    GEMINI_API_KEY =  os.getenv("GEMINI_API_2")
    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel("gemini-1.5-pro")

    """Analyzes any official document and creates a summary Word document report."""

    if not os.path.exists(image_path):
        speak("Sorry Sir, the specified image file does not exist.")
        return

    try:
        speak("Analyzing the official document provided.")
        img = Image.open(image_path).convert('RGB')

        # Generate response from Gemini
        response = model.generate_content([official_doc_prompt, img])

        if response and response.text:
            summary = response.text.strip()

            # Create a Word Document
            doc = Document()
            doc.add_heading('📄 Official Document Summary', 0)

            try:
                doc.add_picture(image_path, width=Inches(5.5))
                doc.add_paragraph("(Original Document Image)").alignment = 1
            except:
                doc.add_paragraph("(⚠️ Could not attach the image)")

            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            doc.add_paragraph(f"📅 Date of Analysis: {timestamp}")
            doc.add_paragraph(summary)

            filename = f"Official_Document_Extraction_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            filepath = os.path.join(os.getcwd(), filename)
            doc.save(filepath)

            speak("Sir, the document summary has been successfully generated and saved in our directory.")
            print(f"Summary saved as {filename}")
            speak("Attempting to open the report of the document.")
            os.startfile(filepath)
            speak("Document opened successfully.")
            return filepath
        else:
            speak("Sorry Sir, I was unable to analyze the document.")
            return

    except Exception as e:
        print(f"Error: {e}")
        speak("An error occurred while analyzing the official document.")
        return

if __name__ == "__main__":
    path = input("Enter path to the official document image: ")
    analyze_official_document(path)
